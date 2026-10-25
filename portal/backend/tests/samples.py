"""Образцы файлов для тестов: строятся в памяти, без файлов в репозитории."""

import io
import struct
import zipfile
import zlib

import docx
import pypdfium2 as pdfium
from PIL import Image

TEXT_LAYER = "This page has a real text layer inside"


def png(width: int = 40, height: int = 30, color: str = "white") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, "PNG")
    return buffer.getvalue()


def png_with_huge_text_chunk(width: int = 300, height: int = 300) -> bytes:
    """PNG, текстовый блок которого распаковывается сверх предела Pillow."""
    source = png(width, height)
    body = b"Comment\x00\x00" + zlib.compress(b"x" * 3_000_000)
    chunk = struct.pack(">I", len(body)) + b"zTXt" + body
    chunk += struct.pack(">I", zlib.crc32(b"zTXt" + body))
    header_end = 8 + 25  # сигнатура и блок IHDR
    return source[:header_end] + chunk + source[header_end:]


def jpeg(width: int = 40, height: int = 30) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, "JPEG")
    return buffer.getvalue()


def scan_pdf(pages: int) -> bytes:
    """PDF без текстового слоя: каждая страница — картинка."""
    images = [Image.new("RGB", (120, 160), "white") for _ in range(pages)]
    buffer = io.BytesIO()
    images[0].save(buffer, "PDF", save_all=True, append_images=images[1:])
    return buffer.getvalue()


def text_pdf(pages: list[str], page_size: tuple[int, int] = (612, 792)) -> bytes:
    """PDF с текстовым слоем; пустая строка — страница без текста (скан)."""
    width, height = page_size
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(len(pages)))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 {height - 72} Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
                f"/Contents {5 + 2 * index} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
            ).encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def docx_file(*paragraphs: str) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_with_pictures(*pictures: bytes) -> bytes:
    """DOCX, в тексте которого по порядку стоят картинки."""
    document = docx.Document()
    document.add_paragraph("Текст перед картинками.")
    for picture in pictures:
        document.add_picture(io.BytesIO(picture))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_DOCX_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)
_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
_REL_TYPES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def paragraph(text: str) -> str:
    """Абзац разметки DOCX для `docx_package`."""
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def docx_package(
    body: str,
    *,
    relations: dict[str, str] | None = None,
    parts: dict[str, bytes] | None = None,
    main: str = "word/document.xml",
    stored: bool = False,
) -> bytes:
    """DOCX, собранный вручную, — для частей пакета, которых python-docx не пишет.

    `body` — содержимое `w:body` основной части `main`; `relations` — её связи-картинки
    (идентификатор → цель; цель `http…` — внешняя), `None` — файла связей нет; `parts` —
    остальные части архива; `stored` — без сжатия.
    """
    links = "".join(
        f'<Relationship Id="{ref}" Type="{_REL_TYPES}/image" Target="{target}"'
        + (' TargetMode="External"/>' if target.startswith("http") else "/>")
        for ref, target in (relations or {}).items()
    )
    folder, _, name = main.rpartition("/")
    files = {
        "[Content_Types].xml": b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
        b'content-types"/>',
        "_rels/.rels": (
            f'<Relationships xmlns="{_RELS}"><Relationship Id="rId1" '
            f'Type="{_REL_TYPES}/officeDocument" Target="{main}"/></Relationships>'
        ).encode(),
        main: f"<w:document {_DOCX_NAMESPACES}><w:body>{body}</w:body></w:document>".encode(),
        **({} if relations is None else {f"{folder}/_rels/{name}.rels": (
            f'<Relationships xmlns="{_RELS}">{links}</Relationships>'.encode()
        )}),
        **(parts or {}),
    }  # fmt: skip
    buffer = io.BytesIO()
    method = zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED
    with zipfile.ZipFile(buffer, "w", method) as archive:
        for part_name, content in files.items():
            archive.writestr(part_name, content)
    return buffer.getvalue()


def zip_claiming_size(entries: dict[str, bytes], method: int, name: str, size: int) -> bytes:
    """ZIP из `entries` со сжатием `method`, где у записи `name` заявлен размер `size`.

    Размер подменяется и в локальном заголовке, и в оглавлении: так архив-бомба выглядит
    по оглавлению маленьким.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", method) as archive:
        for part_name, content in entries.items():
            archive.writestr(part_name, content)
    data = bytearray(buffer.getvalue())
    encoded = name.encode()
    for signature, size_at, name_at in ((b"PK\x03\x04", 22, 30), (b"PK\x01\x02", 24, 46)):
        start = 0
        while (start := data.find(signature, start)) >= 0:
            if data[start + name_at : start + name_at + len(encoded)] == encoded:
                data[start + size_at : start + size_at + 4] = size.to_bytes(4, "little")
            start += 4
    return bytes(data)


def zip_with_method(source: bytes, name: str, method: int) -> bytes:
    """Тот же ZIP, где запись `name` сжата способом `method` (bzip2, LZMA), прочие — как были."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(source)) as original,
        zipfile.ZipFile(buffer, "w") as archive,
    ):
        for info in original.infolist():
            kind = method if info.filename == name else info.compress_type
            archive.writestr(info.filename, original.read(info), compress_type=kind)
    return buffer.getvalue()


def zip_directory_claims(source: bytes, size: int, *, zip64: bool = False) -> bytes:
    """ZIP, запись конца которого заявляет одну запись и оглавление размером `size`.

    `zip64` — размер заявлен в записи ZIP64 EOCD, а обычная запись конца честная.
    """
    end = source.rindex(b"PK\x05\x06")
    record = bytearray(source[end : end + 22])
    struct.pack_into("<HH", record, 8, 1, 1)
    if not zip64:
        struct.pack_into("<I", record, 12, size)
        return source[:end] + bytes(record)
    offset = struct.unpack_from("<I", record, 16)[0]
    zip64_end = struct.pack("<4sQHHLLQQQQ", b"PK\x06\x06", 44, 45, 45, 0, 0, 1, 1, size, offset)
    locator = struct.pack("<4sLQL", b"PK\x06\x07", 0, end, 1)
    return source[:end] + zip64_end + locator + bytes(record)


def pdf_with_picture(picture_jpeg: bytes, text: str = TEXT_LAYER, copies: int = 1) -> bytes:
    """PDF, на странице которого есть и текстовый слой, и встроенная картинка JPEG.

    `copies` — сколько раз картинка стоит на странице.
    """
    document = pdfium.PdfDocument(text_pdf([text]))
    page = document[0]
    for _ in range(copies):
        image = pdfium.PdfImage.new(document)
        image.load_jpeg(io.BytesIO(picture_jpeg), inline=True)
        image.set_matrix(pdfium.PdfMatrix().scale(200, 150))
        page.insert_obj(image)
    page.gen_content()
    buffer = io.BytesIO()
    document.save(buffer)
    page.close()
    document.close()
    return buffer.getvalue()
