"""Экспорт ответа или диалога в DOCX на `python-docx` (docs/portal-api.md §5.9).

Текст ответов модели и документов вставляется только как данные — текстом абзацев и
ячеек. Гиперссылок, полей и внешних связей в файле нет: при открытии он ничего не
подтягивает. HTML из текста не интерпретируется и остаётся обычными символами.
"""

import io
import re
from collections.abc import Sequence

import docx
from docx.document import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run

from portal.dialogs.domain import Dialog, Docparse
from portal.dialogs.footnotes import cited_numbers
from portal.dialogs.markdown import (
    Code,
    Heading,
    ListBlock,
    Paragraph,
    Table,
    parse_blocks,
    parse_spans,
)
from portal.dialogs.service import MessageView

_CODE_FONT = "Courier New"
_CODE_SIZE = Pt(9)
_MAX_HEADING_LEVEL = 4
_TABLE_STYLE = "Table Grid"
# Всё, кроме допустимого в XML 1.0: \\t, \\n, \\r и символы от пробела, без суррогатов.
_XML_FORBIDDEN = re.compile("[^\\t\\n\\r\\x20-\\ud7ff\\ue000-\\ufffd\\U00010000-\\U0010ffff]")


def xml_safe(text: str) -> str:
    """Убрать символы, недопустимые в XML 1.0: с ними файл не собрать.

    Вертикальная табуляция, перевод страницы и другие управляющие символы обычны для
    текста из Word, PDF и распознанных сканов.
    """
    return _XML_FORBIDDEN.sub("", text)


# Весь текст попадает в документ только через эти четыре функции.


def _add_run(paragraph: DocxParagraph, text: str) -> Run:
    return paragraph.add_run(xml_safe(text))


def _add_paragraph(document: Document, text: str = "", style: str | None = None) -> DocxParagraph:
    paragraph = document.add_paragraph(style=style)
    if text:
        _add_run(paragraph, text)
    return paragraph


def _add_heading(document: Document, text: str, level: int) -> None:
    document.add_heading(xml_safe(text), level=level)


def _add_table(document: Document, rows: Sequence[Sequence[str]], *, markup: bool) -> None:
    """Таблица с заголовком в первой строке; `markup` — разбирать ли оформление в ячейках.

    Ячейки строки берутся один раз: обращение к ячейке по номерам строки и столбца в
    `python-docx` перебирает всю таблицу, и большая таблица строилась бы минутами.
    """
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = _TABLE_STYLE
    for row_index, (table_row, values) in enumerate(zip(table.rows, rows, strict=True)):
        for cell, value in zip(table_row.cells, values, strict=False):
            paragraph = cell.paragraphs[0]
            if markup:
                _write_spans(paragraph, value)
            else:
                _add_run(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def _write_spans(paragraph: DocxParagraph, text: str) -> None:
    for span in parse_spans(text):
        run = _add_run(paragraph, span.text)
        run.bold = span.bold or None
        run.italic = span.italic or None
        if span.code:
            run.font.name = _CODE_FONT


def write_markdown(document: Document, text: str, heading_offset: int = 0) -> None:
    """Перенести Markdown в документ: заголовки, абзацы, списки, таблицы, блоки кода."""
    for block in parse_blocks(text):
        if isinstance(block, Heading):
            level = min(block.level + heading_offset, _MAX_HEADING_LEVEL)
            _add_heading(document, block.text, level)
        elif isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _write_spans(_add_paragraph(document, style=style), item)
        elif isinstance(block, Table):
            _add_table(document, block.rows, markup=True)
        elif isinstance(block, Code):
            for line in block.text.split("\n"):
                paragraph = _add_paragraph(document)
                paragraph.paragraph_format.space_after = Pt(0)
                run = _add_run(paragraph, line)
                run.font.name = _CODE_FONT
                run.font.size = _CODE_SIZE
            _add_paragraph(document)
        elif isinstance(block, Paragraph):
            _write_spans(_add_paragraph(document), block.text)


def _write_label(document: Document, label: str) -> None:
    _add_run(_add_paragraph(document), label).bold = True


def _write_answer(document: Document, view: MessageView, heading_offset: int) -> None:
    message = view.message
    write_markdown(document, message.content, heading_offset)
    sources = message.sources or []
    # Тем же правилом сносок, что при сохранении ответа: только упомянутые вне кода.
    cited = cited_numbers(message.content, [int(source["n"]) for source in sources])
    listed = [source for source in sources if int(source["n"]) in cited]
    if listed:
        _write_label(document, "Источники")
        for source in listed:
            page = f", стр. {source['page']}" if source.get("page") is not None else ""
            _add_paragraph(document, f"[{source['n']}] {source['document_title']}{page}")


def _write_docparse(document: Document, docparse: Docparse) -> None:
    _add_paragraph(document, f"Файл: {docparse.file_name}. Шаблон: {docparse.template_title}.")
    if docparse.fields:
        rows = [("Реквизит", "Значение")] + [
            (str(item["title"]), str(item["value"] or "—")) for item in docparse.fields
        ]
        # Значения реквизитов — данные документа: в ячейки они идут без разбора разметки.
        _add_table(document, rows, markup=False)
    if docparse.summary:
        _add_heading(document, "Краткое содержание", level=2)
        write_markdown(document, docparse.summary, heading_offset=2)


class DocxExporter:
    """Собирает файл DOCX; работа процессорная — вызывается в пуле потоков."""

    def answer(self, title: str, view: MessageView) -> bytes:
        """Один ответ."""
        document = docx.Document()
        _add_heading(document, title, level=1)
        _write_answer(document, view, heading_offset=1)
        return _saved(document)

    def dialog(
        self, title: str, dialog: Dialog, docparse: Docparse | None, views: Sequence[MessageView]
    ) -> bytes:
        """Диалог целиком: вопросы и ответы с пометками «Вы» и «Ответ»."""
        document = docx.Document()
        _add_heading(document, title, level=1)
        if docparse is not None:
            _write_docparse(document, docparse)
        for view in views:
            message = view.message
            if message.role == "user":
                _write_label(document, "Вы")
                if message.content:
                    # Вопрос пользователя — как введён, без разбора разметки.
                    for line in message.content.split("\n"):
                        _add_paragraph(document, line)
                if view.attachments:
                    names = ", ".join(item.file_name for item in view.attachments)
                    _add_paragraph(document, f"Вложения: {names}")
            else:
                _write_label(document, "Ответ")
                _write_answer(document, view, heading_offset=1)
        return _saved(document)


def _saved(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
